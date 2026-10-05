"""Local Eastmoney file preview, explicit mapping, atomic deduplicated import."""
import base64,csv,hashlib,io,json,re,uuid,zipfile
from datetime import datetime
from decimal import DecimalException
from pathlib import Path
from .domain import Security,dec
from .providers import ProviderError
from .storage import utcnow

ALIASES={
 'ticker':['证券代码','股票代码','标的代码','代码'], 'name':['证券名称','股票名称','名称'],
 'exchange':['市场','证券市场','交易市场','交易所'], 'quantity':['股票余额','证券数量','持仓数量','持仓','数量','成交数量'],
 'broker_cost':['成本价','成本价格','成本'], 'broker_value':['市值','证券市值','参考市值'],
 'broker_gain':['持仓盈亏','浮动盈亏','参考盈亏'], 'day':['成交日期','交易日期','发生日期','日期'],
 'kind':['业务名称','业务类型','交易类别','操作','买卖标志'], 'amount_cny':['人民币交收金额','人民币发生金额','发生金额','清算金额'],
 'fee_cny':['手续费合计','费用合计','费用'], 'broker_ref':['流水号','合同编号','成交编号'],
 'currency':['币种','货币'], 'cost_currency':['成本币种'],
}
KINDS={'证券买入':'BUY','买入':'BUY','证券卖出':'SELL','卖出':'SELL','股息入账':'DIVIDEND','红利入账':'DIVIDEND','分红':'DIVIDEND','利息归本':'INTEREST','利息':'INTEREST','银行转证券':'DEPOSIT','银行转存':'DEPOSIT','入金':'DEPOSIT','证券转银行':'WITHDRAWAL','银行转取':'WITHDRAWAL','出金':'WITHDRAWAL','费用':'FEE','送股':'SPLIT','股份增加':'SPLIT','股份减少':'SPLIT','配股缴款':'BUY'}

def number(value,nullable=False):
    s=str('' if value is None else value).strip().replace(',','').replace('，','').replace('￥','').replace('¥','')
    if '%' in s:raise ValueError('金额或数量字段不能使用百分比')
    if not s or s in ('--','-'):
        if nullable:return None
        raise ValueError('缺少数字')
    if s.endswith('万'):return str(dec(s[:-1])*10000)
    result=dec(s)
    if abs(result)>10**12:raise ValueError('数字超过账户导入支持范围，请核对单位')
    return str(result)

def day(value):
    if isinstance(value,datetime):return value.date().isoformat()
    s=str(value).strip().replace('/','-')
    if re.fullmatch(r'\d{8}',s):s=s[:4]+'-'+s[4:6]+'-'+s[6:]
    return datetime.fromisoformat(s).date().isoformat()

def preview(store,body):
    filename=Path(str(body.get('filename','导入.csv'))).name
    try:blob=base64.b64decode(body['content_base64'],validate=True)
    except Exception:raise ProviderError('文件编码无效，请重新选择导出文件。')
    if len(blob)>10*1024*1024:raise ProviderError('单个导入文件最多10MB，请按月份导出。')
    try:
        if filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                if sum(i.file_size for i in z.infolist())>80*1024*1024:raise ValueError('展开文件过大')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
            try:raw=list(book.active.iter_rows(values_only=True))
            finally:book.close()
        else:
            try:text=blob.decode('utf-8-sig')
            except UnicodeDecodeError:text=blob.decode('gb18030')
            dialect=csv.Sniffer().sniff(text[:8000],delimiters=',\t;')
            raw=list(csv.reader(io.StringIO(text),dialect))
        if len(raw)>20001:raise ValueError('超过20000条记录')
        header_index=next(i for i,r in enumerate(raw[:30]) if any(str(x).strip() in ALIASES['ticker']+ALIASES['day'] for x in r))
        headers=[str(x or '').strip() for x in raw[header_index]]
        if len(set(headers))!=len(headers):raise ValueError('表头有重复或空白列，请去掉空白列')
        mapping={k:next((x for x in v if x in headers),None) for k,v in ALIASES.items()}
        for k,v in body.get('mapping',{}).items():
            if k in ALIASES and (v in headers or v==''):mapping[k]=v or None
        import_kind=body.get('kind') or ('transactions' if mapping['kind'] else 'holdings')
        if import_kind not in ('holdings','transactions'):raise ValueError('类型无效')
        parsed=[];issues=[]
        for n,values in enumerate(raw[header_index+1:],header_index+2):
            if not any(x is not None and str(x).strip() for x in values):continue
            obj=dict(zip(headers,values))
            get=lambda k:obj.get(mapping.get(k))
            try:
                row={'row':n,'raw':{str(k):str(v or '') for k,v in obj.items()}}
                code=str(get('ticker') or '').strip().removesuffix('.0')
                if code:
                    market=str(get('exchange') or '')
                    ex='HK' if any(x in market for x in ('港','HK')) or len(code)<=5 else ('SH' if code.startswith(('5','6','9')) else 'BJ' if code.startswith(('4','8','92')) else 'SZ')
                    s=Security(exchange=ex,ticker=code,name=str(get('name') or code)).model_dump()
                    row.update(s,security_id=ex+':'+s['ticker'])
                if import_kind=='holdings':
                    if not code:raise ValueError('持仓缺少证券代码')
                    row.update(quantity=number(get('quantity')),broker_cost=number(get('broker_cost'),True),broker_value=number(get('broker_value'),True),broker_gain=number(get('broker_gain'),True),cost_currency=str(get('cost_currency') or body.get('cost_currency','UNKNOWN')).upper())
                    if dec(row['quantity'])<0:raise ValueError('不支持负数量持仓')
                    row['cost_currency']={'人民币':'CNY','RMB':'CNY','港币':'HKD','港元':'HKD'}.get(row['cost_currency'],row['cost_currency'])
                    if row['cost_currency'] not in ('CNY','HKD','UNKNOWN'):raise ValueError('成本币种无法识别')
                    row['as_of']=body.get('as_of') or utcnow()
                else:
                    action=str(get('kind') or '').strip()
                    kind=action if action in set(KINDS.values()) else KINDS.get(action)
                    if not kind:raise ValueError('未识别的业务类型：'+action)
                    if kind in ('BUY','SELL','SPLIT') and not code:raise ValueError('交易缺少证券代码')
                    currency=str(get('currency') or body.get('settlement_currency','UNKNOWN')).upper()
                    if currency not in ('CNY','人民币','RMB'):raise ValueError('请映射人民币交收金额并确认结算币种；不能把港元成交额当成人民币')
                    quantity=dec(number(get('quantity'),True) or '0')
                    if kind!='SPLIT':quantity=abs(quantity)
                    elif action=='股份减少':quantity=-abs(quantity)
                    if (kind in ('BUY','SELL') and quantity<=0) or (kind=='SPLIT' and quantity==0):raise ValueError('成交数量必须大于零')
                    row.update(kind=kind,day=day(get('day')),quantity=str(quantity),amount_cny=str(abs(dec(number(get('amount_cny'))))),fee_cny=str(abs(dec(number(get('fee_cny'),True) or '0'))),broker_ref=str(get('broker_ref') or ''),amount_includes_fees=bool(body.get('amount_includes_fees',True)))
                    core=[row.get(k) for k in ('security_id','kind','day','quantity','amount_cny','fee_cny','amount_includes_fees')]
                    row['dedupe_key']='ref:'+row['broker_ref'] if row['broker_ref'] else 'value:'+hashlib.sha256(json.dumps(core).encode()).hexdigest()
                    existing=store.rows('SELECT id FROM ledger WHERE account_id=? AND dedupe_key=?',('eastmoney',row['dedupe_key']))
                    row['duplicate']='exact' if existing and row['broker_ref'] else 'suspected' if existing or any(x.get('dedupe_key')==row['dedupe_key'] for x in parsed) else None
                parsed.append(row)
            except (ValueError,TypeError,DecimalException) as e:issues.append({'row':n,'message':str(e)})
    except ProviderError:raise
    except Exception as e:raise ProviderError('无法识别导出文件：'+str(e)) from e
    iid=uuid.uuid4().hex
    result={'id':iid,'filename':filename,'kind':import_kind,'headers':headers,'mapping':mapping,'rows':parsed,'issues':issues,'status':'preview','account_id':'eastmoney','replace_snapshot':body.get('replace_snapshot') is True}
    store.execute('INSERT INTO imports VALUES(?,?,?,?,?,?)',(iid,filename,hashlib.sha256(blob).hexdigest(),json.dumps(result,ensure_ascii=False),'preview',utcnow()))
    return result

def commit(store,iid,body):
    selected=set(body.get('exclude_rows',[]))
    with store.connect() as db:
        db.execute('BEGIN IMMEDIATE')
        batch=db.execute('SELECT * FROM imports WHERE id=?',(iid,)).fetchone()
        if not batch:raise ProviderError('导入预览不存在。')
        if batch['status']=='committed':return {'ok':True,'already_committed':True}
        data=json.loads(batch['payload'])
        if data['issues'] and not body.get('accept_valid_rows'):raise ProviderError('请核对异常行，或明确选择仅导入有效记录。')
        if data.get('replace_snapshot') and (data['issues'] or selected):raise ProviderError('完整快照替换需核对全部记录；有异常或排除行时请先修正，或改用合并导入。')
        if data.get('replace_snapshot') and data['kind']=='holdings':
            if not data['rows']:raise ProviderError('空文件不能替换当前持仓。')
            db.execute("DELETE FROM holding_snapshots WHERE account_id='eastmoney'")
            db.execute('DELETE FROM positions')
        count=0
        for row in data['rows']:
            if row['row'] in selected:continue
            sid=row.get('security_id')
            if sid:
                db.execute('INSERT INTO securities(id,exchange,ticker,name,currency) VALUES(?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name',(sid,row['exchange'],row['ticker'],row['name'],'HKD' if row['exchange']=='HK' else 'CNY'))
            if data['kind']=='holdings':
                db.execute('INSERT OR REPLACE INTO holding_snapshots VALUES(?,?,?,?,?,?,?,?)',('eastmoney',sid,row['quantity'],row['broker_cost'],row['cost_currency'],row['broker_value'],row['broker_gain'],row['as_of']))
                if dec(row['quantity'])==0:db.execute('DELETE FROM positions WHERE security_id=?',(sid,))
                else:db.execute('INSERT OR REPLACE INTO positions VALUES(?,?,?,?,?)',(sid,row['quantity'],row['broker_cost'] or '0','东方财富快照；券商成本独立保存',utcnow()))
                count+=1
            else:
                existing=db.execute('SELECT payload FROM ledger WHERE account_id=? AND dedupe_key=?',('eastmoney',row['dedupe_key'])).fetchone()
                if existing:
                    if row['broker_ref']:
                        old=json.loads(existing['payload'])
                        if any(old.get(k)!=row.get(k) for k in ('kind','day','quantity','amount_cny','fee_cny','security_id','amount_includes_fees')):raise ProviderError('相同流水号内容冲突，请先核对。')
                        continue
                    raise ProviderError('发现疑似重复交易，请在预览中排除，不能静默丢弃。')
                db.execute('INSERT INTO ledger(account_id,security_id,kind,day,quantity,amount_cny,fee_cny,broker_ref,dedupe_key,payload) VALUES(?,?,?,?,?,?,?,?,?,?)',('eastmoney',sid,row['kind'],row['day'],row['quantity'],row['amount_cny'],row['fee_cny'],row['broker_ref'],row['dedupe_key'],json.dumps(row,ensure_ascii=False)))
                count+=1
        if count:
            db.execute("UPDATE accounts SET mode='snapshot',history_complete=0 WHERE id='eastmoney'")
        db.execute("UPDATE imports SET status='committed' WHERE id=?",(iid,))
    return {'ok':True,'imported':count,'note':'流水已保存，账户仍按快照展示；完整历史核对后才切换核算模式。'}
